"""Данные для макета дашборда: метки, прогнозы попыток, справочник остановок.

Запуск: .venv/bin/python web/scripts/mock_data.py
Пишет web/public/data/*.json. Сырые данные в git не идут, собранное — небольшое.
"""
import json
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "raw"
EXT = ROOT / "data" / "external"
SUB = ROOT / "artifacts" / "submissions"
OUT = ROOT / "web" / "public" / "data"
ROUTES = [1, 5, 7, 11, 12, 17, 25, 26, 28, 50]
MAIN_MODEL = "2026-09-27_v54_t1_sources"  # итоговая попытка, 0,90715


def load_labels():
    parts = [pd.read_csv(RAW / "labels" / f, sep=";") for f in ("labels_day_train.csv", "labels_day_test.csv")]
    df = pd.concat(parts)
    df["date"] = pd.to_datetime(df["date"])
    return df


def grid(start, end):
    days = pd.date_range(start, end, freq="D")
    idx = pd.MultiIndex.from_product([ROUTES, days, range(24)], names=["route", "date", "hour"])
    return idx


def to_matrix(series, start, end):
    """route → [день][час] за период, пустой час — 0."""
    full = series.reindex(grid(start, end), fill_value=0)
    out = {}
    for r in ROUTES:
        out[str(r)] = full.loc[r].to_numpy().reshape(-1, 24).round().astype(int).tolist()
    return out


def profile(df, start, end):
    part = df[(df.date >= start) & (df.date <= end)].copy()
    full = part.set_index(["route", "date", "hour"])["boardings"].reindex(grid(start, end), fill_value=0).reset_index()
    full["dow"] = full.date.dt.dayofweek
    return full.groupby(["route", "dow", "hour"])["boardings"]


def apply_profile(prof, start, end):
    g = grid(start, end).to_frame(index=False)
    g["dow"] = g.date.dt.dayofweek
    g = g.merge(prof.rename("p").reset_index(), on=["route", "dow", "hour"], how="left").fillna(0)
    return g.set_index(["route", "date", "hour"])["p"]


def wape_score(y, p):
    return round(max(0.0, 1 - float(np.abs(y - p).sum() / y.sum())), 4)


def stops():
    wb = openpyxl.load_workbook(RAW / "spravochniki" / "Хакатон_справочники_трамвай_10_маршрутов.xlsx", read_only=True)
    sheet = lambda prefix: next(ws for ws in wb.worksheets if ws.title.startswith(prefix))
    routes_rows = list(sheet("Маршруты").iter_rows(values_only=True))[2:]
    seq = list(sheet("Порядок_с_координатами").iter_rows(values_only=True))[1:]
    names = {}
    for r in routes_rows:
        if r[2] and r[2].isdigit():
            names[int(r[2])] = {"name": r[3], "since": r[8]}
    by_route = {}
    for row in seq:
        route, trip, direction = row[1], row[4], row[6]
        if not route or not route.isdigit() or str(direction) != "0":
            continue
        by_route.setdefault(int(route), {}).setdefault(trip, []).append(
            (int(row[9]), row[10], row[14], float(row[15]), float(row[16]))
        )
    out = {}
    for route, trips in by_route.items():
        longest = max(trips.values(), key=len)
        longest.sort()
        seen, pts = set(), []
        for _, sid, name, lat, lon in longest:
            if sid in seen:
                continue
            seen.add(sid)
            pts.append({"id": sid, "name": name, "lat": round(lat, 6), "lon": round(lon, 6)})
        out[route] = pts
    return names, out


def mos_stops(route):
    """Трасса по открытым данным Москвы: у набора нет порядка, он восстанавливается по близости."""
    path = EXT / "mos_busstops_752.json"
    if not path.exists():
        return []
    import re
    groups = {}
    for row in json.loads(path.read_text()):
        if str(route) not in re.findall(r"Тм\s*(\d+)", row["RouteNumbers"] or ""):
            continue
        g = groups.setdefault(row["StationName"], [])
        g.append((float(row["Latitude_WGS84"]), float(row["Longitude_WGS84"]), str(row["ID"])))
    pts = [
        {"id": v[0][2], "name": name, "lat": round(sum(p[0] for p in v) / len(v), 6), "lon": round(sum(p[1] for p in v) / len(v), 6)}
        for name, v in groups.items()
    ]
    if len(pts) < 2:
        return pts
    d = lambda a, b: ((a["lat"] - b["lat"]) ** 2 + ((a["lon"] - b["lon"]) * 0.56) ** 2) ** 0.5
    cy = sum(p["lat"] for p in pts) / len(pts)
    cx = sum(p["lon"] for p in pts) / len(pts)
    start = max(pts, key=lambda p: d(p, {"lat": cy, "lon": cx}))
    chain, rest = [start], [p for p in pts if p is not start]
    while rest:
        nxt = min(rest, key=lambda p: d(chain[-1], p))
        chain.append(nxt)
        rest.remove(nxt)
    return chain


def ours_month(train_end: str, start: str, end: str) -> pd.Series:
    """Ядро основной модели на месяц истории — тот же расчёт, что в журнале и в tramflow.protocol."""
    import sys
    sys.path.insert(0, str(ROOT))
    from tramflow import backtest, protocol, versions
    d = backtest.detail(train_end, start, end, **{**versions.BASE, **protocol.core(MAIN_MODEL.split("_", 1)[1])})
    d["date"] = pd.to_datetime(d["date"])
    return d.set_index(["route", "date", "hour"])["prediction"]


def registered() -> set:
    """Версии из реестра tramflow.versions: на экран идут только они."""
    import sys
    sys.path.insert(0, str(ROOT))
    from tramflow import versions
    return set(versions.VERSIONS)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    df = load_labels()
    fact = df.set_index(["route", "date", "hour"])["boardings"]

    # Факт сентября—октября и прогноз на них честным бэктестом ядра основной модели.
    fact_so = to_matrix(fact, "2025-09-01", "2025-10-31")
    sep_fc = ours_month("2025-08-31", "2025-09-01", "2025-09-30")
    oct_fc = ours_month("2025-09-30", "2025-10-01", "2025-10-31")
    prof_sep = apply_profile(profile(df, "2025-09-01", "2025-09-30").median(), "2025-10-01", "2025-10-31")
    fc_so = to_matrix(pd.concat([sep_fc, oct_fc]), "2025-09-01", "2025-10-31")

    # Ноябрь—декабрь: попытки из журнала.
    known = registered()
    # первые три попытки собраны до реестра версий, остальные берутся только из него
    # в журнале есть записи без даты в имени — это пробы вне реестра, на экран они не идут
    dated = [a for a in json.loads((SUB / "log.json").read_text()) if a["name"][:4].isdigit()]
    log = [a for a in dated
           if a["name"].split("_", 1)[1] in known or int(a["name"].split("_")[1][1:]) < 4]
    models_nd = {}
    for item in log:
        path = SUB / f"{item['name']}.csv"
        if not path.exists():
            continue
        s = pd.read_csv(path, sep=";")
        s["date"] = pd.to_datetime(s["date"])
        models_nd[item["name"]] = to_matrix(s.set_index(["route", "date", "hour"])["prediction"], "2025-11-01", "2025-12-31")

    # Интервал: квантили 10 и 90 % к медиане по «маршрут × день недели × час» за сентябрь—октябрь.
    g = profile(df, "2025-09-01", "2025-10-31")
    med, lo, hi = g.median(), g.quantile(0.1), g.quantile(0.9)
    band = {}
    for r in ROUTES:
        m = med.loc[r].to_numpy().reshape(7, 24)
        with np.errstate(divide="ignore", invalid="ignore"):
            band[str(r)] = {
                "lo": np.nan_to_num(lo.loc[r].to_numpy().reshape(7, 24) / m, nan=1, posinf=1).clip(0, 3).round(3).tolist(),
                "hi": np.nan_to_num(hi.loc[r].to_numpy().reshape(7, 24) / m, nan=1, posinf=1).clip(0, 3).round(3).tolist(),
            }
    # у маршрута 5 истории нет, прогноз — профиль 25-го, и разброс берётся у него же
    band["5"] = band["25"]

    # Сравнение решений на октябре.
    y = fact.reindex(grid("2025-10-01", "2025-10-31"), fill_value=0)
    mean_route = df[df.date <= "2025-08-31"].groupby("route")["boardings"].mean()
    candidates = {
        "baseline": ("Среднее маршрута за все часы", "аналог выданного baseline",
                     pd.Series(y.index.get_level_values(0).map(mean_route).fillna(0).to_numpy(), index=y.index)),
        "prof_aug": ("Профиль, медиана за август", "маршрут × день недели × час",
                     apply_profile(profile(df, "2025-08-01", "2025-08-31").median(), "2025-10-01", "2025-10-31")),
        "prof_jan_aug": ("Профиль, медиана за январь—август", "маршрут × день недели × час",
                         apply_profile(profile(df, "2025-01-01", "2025-08-31").median(), "2025-10-01", "2025-10-31")),
        "prof_sep": ("Профиль, медиана за сентябрь", "маршрут × день недели × час", prof_sep),
        "ours": ("TramFlow", MAIN_MODEL.split("_", 1)[1], oct_fc),
    }
    compare = []
    oct_series = {"fact": to_matrix(y, "2025-10-01", "2025-10-31")}
    wd = y.index.get_level_values(1).dayofweek < 5
    for key, (title, note, pred) in candidates.items():
        pred = pred.reindex(y.index, fill_value=0)
        compare.append({
            "key": key, "title": title, "note": note,
            "oct": wape_score(y.to_numpy(), pred.to_numpy()),
            "oct_weekday": wape_score(y[wd].to_numpy(), pred[wd].to_numpy()),
            "oct_weekend": wape_score(y[~wd].to_numpy(), pred[~wd].to_numpy()),
            "bias_pct": round(float((pred.sum() - y.sum()) / y.sum() * 100), 1),
        })
        oct_series[key] = to_matrix(pred, "2025-10-01", "2025-10-31")

    # История по дням для горизонта «год».
    daily = df.groupby(["route", "date"])["boardings"].sum()
    days = pd.date_range("2025-01-01", "2025-10-31")
    history = {str(r): daily.reindex(pd.MultiIndex.from_product([[r], days]), fill_value=0).astype(int).tolist() for r in ROUTES}

    cal = pd.read_csv(EXT / "isdayoff_2025.csv")
    w = pd.read_csv(EXT / "open_meteo_2025-01-01_2025-12-31.csv", parse_dates=["time"])
    w["date"] = w.time.dt.strftime("%Y-%m-%d")
    wd_ = w.groupby("date").agg(t=("temperature_2m", "mean"), rain=("rain", "sum"), snow=("snowfall", "sum"))
    weather = {d: [round(r.t, 1), round(r.rain, 1), round(r.snow, 1)] for d, r in wd_.iterrows()}

    names, stop_lists = stops()
    routes = []
    for r in ROUTES:
        own = stop_lists.get(r, [])
        stops_ = own or mos_stops(r)
        routes.append({
            "id": str(r),
            "name": names.get(r, {}).get("name") or (f"{stops_[0]['name']} — {stops_[-1]['name']}" if stops_ else f"Маршрут {r}"),
            "since": names.get(r, {}).get("since"),
            "stops": stops_,
            "source": "org" if own else ("mos" if stops_ else None),
        })

    import re
    for r in routes:
        r["name"] = re.sub(r'"([^"]+)"', r"«\1»", r["name"]).replace(" - ", " — ")
        for st in r["stops"]:
            st["name"] = re.sub(r'"([^"]+)"', r"«\1»", st["name"])
    dump = lambda name, obj: (OUT / name).write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")))
    dump("routes.json", routes)
    dump("series.json", {"start": "2025-09-01", "fact": fact_so, "backtest": fc_so, "band": band})
    for a in log:
        if isinstance(a.get("score"), str):
            a["score"] = float(a["score"])
    dump("models.json", {"main": MAIN_MODEL, "log": log, "compare": compare, "nd": models_nd, "oct": oct_series})
    dump("history.json", {"start": "2025-01-01", "daily": history})
    dump("context.json", {"calendar": dict(zip(cal.date, cal.code.astype(int))), "weather": weather})
    for f in sorted(OUT.glob("*.json")):
        print(f.name, f.stat().st_size // 1024, "КБ")
    print({c["key"]: c["oct"] for c in compare})


if __name__ == "__main__":
    main()
